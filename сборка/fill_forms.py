# -*- coding: utf-8 -*-
"""Пять заполненных официальных форм самооценки, по одной на участника."""
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent if Path(__file__).resolve().parent.name == "сборка" else Path(__file__).resolve().parent
TEMPLATE = ROOT.parent / "Форма самооценки команды.docx"
OUT = ROOT / "самооценка"

PEERS = [
    ("Подъячий Глеб Русланович", "организация, Python, выводы"),
    ("Керимов Андрей Эдуардович", "Java, два алгоритма Arrays.sort"),
    ("Мамбетов Родион Рудемович", "C# Array.Sort и OrderBy"),
    ("Черакаев Антон Александрович", "C++ std::sort и stable_sort"),
    ("Кулиш Артём Евгеньевич", "JavaScript и демонстрация"),
]

PEOPLE = [
    {
        "file": "Подъячий_Глеб.docx",
        "name": "Подъячий Глеб Русланович",
        "roles": ["Лидер команды", "Исследователь", "Докладчик"],
        "tasks": [
            ("Распределил пять языков по участникам", "У каждого есть своя часть на 1 минуту"),
            ("Разобрал list.sort в Python 3.14", "Timsort и Powersort, устойчивость подтверждена"),
            ("Собрал итоговые рекомендации", "Когда брать устойчивую сортировку, а когда нет"),
        ],
        "hard": "Нужно было не смешать пять языков в один рассказ и оставить каждому понятный кусок.",
        "best": "Общая таблица: по ней сразу видно, где алгоритм устойчивый.",
        "better": "Добавить замер памяти и отдельный прогон Firefox, чтобы сравнить его слияние с Timsort в V8.",
    },
    {
        "file": "Керимов_Андрей.docx",
        "name": "Керимов Андрей Эдуардович",
        "roles": ["Исследователь", "Разработчик", "Докладчик"],
        "tasks": [
            ("Разобрал Arrays.sort для int[]", "Dual-Pivot Quicksort, 48 мс на случайном миллионе"),
            ("Разобрал Arrays.sort для Integer[]", "Timsort, на том же массиве 260 мс"),
            ("Проверил устойчивость объектов", "Порядок равных элементов сохранился"),
        ],
        "hard": "В документации легко решить, что в Java один алгоритм. На самом деле их два.",
        "best": "Наглядная разница 48 мс и 260 мс на одном и том же миллионе чисел.",
        "better": "Добавить замер памяти и отдельный прогон Firefox, чтобы сравнить его слияние с Timsort в V8.",
    },
    {
        "file": "Мамбетов_Родион.docx",
        "name": "Мамбетов Родион Рудемович",
        "roles": ["Исследователь", "Разработчик", "Докладчик"],
        "tasks": [
            ("Разобрал Array.Sort", "introsort: быстрая сортировка, вставки и пирамида"),
            ("Отделил OrderBy от Array.Sort", "OrderBy устойчивый, Array.Sort нет"),
            ("Проверил устойчивость на 400 элементах", "Программа показала, что порядок сломан"),
        ],
        "hard": "Короткий массив сортируется вставками и может случайно сохранить порядок. Проверку пришлось делать на 400 элементах.",
        "best": "Чёткое правило для защиты: нужен прежний порядок равных записей, вызываем OrderBy.",
        "better": "Добавить замер памяти и отдельный прогон Firefox, чтобы сравнить его слияние с Timsort в V8.",
    },
    {
        "file": "Черакаев_Антон.docx",
        "name": "Черакаев Антон Александрович",
        "roles": ["Исследователь", "Разработчик", "Докладчик"],
        "tasks": [
            ("Разобрал std::sort в MSVC", "introsort, на месте, не устойчивый"),
            ("Сравнил его со std::stable_sort", "Слияние, порядок равных сохраняется"),
            ("Собрал программу bench.cpp с оптимизацией", "Оба вызова замерены на тех же файлах"),
        ],
        "hard": "std::sort на коротком куске выглядит устойчивым. На 400 элементах неустойчивость уже видна.",
        "best": "Две функции стандартной библиотеки закрывают оба случая: быстро и устойчиво.",
        "better": "Добавить замер памяти и отдельный прогон Firefox, чтобы сравнить его слияние с Timsort в V8.",
    },
    {
        "file": "Кулиш_Артём.docx",
        "name": "Кулиш Артём Евгеньевич",
        "roles": ["Исследователь", "Разработчик", "Докладчик"],
        "tasks": [
            ("Разобрал Array.prototype.sort в Node.js 24", "V8, Timsort, сортировка устойчивая"),
            ("Показал ловушку без компаратора", "Числа 10, 2, 1 сортируются как строки"),
            ("Собрал страницу demo.html", "На защите видно и ловушку, и таблицу времени"),
        ],
        "hard": "Стандарт требует устойчивость, но не называет алгоритм. Замер сделан в V8, не во всех браузерах.",
        "best": "Ловушка со строковой сортировкой объясняется за полминуты и сразу запоминается.",
        "better": "Добавить замер памяти и отдельный прогон Firefox, чтобы сравнить его слияние с Timsort в V8.",
    },
]


def replace_text(paragraph, old, new):
    full = "".join(run.text for run in paragraph.runs) or paragraph.text
    if old not in full:
        return False
    updated = full.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = updated
    return True


def set_cell(cell, text):
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
        paragraph.runs[0].font.size = Pt(12)
    else:
        run = paragraph.add_run(text)
        run.font.size = Pt(12)


def fill(person):
    target = OUT / person["file"]
    shutil.copyfile(TEMPLATE, target)
    doc = Document(target)
    header = doc.paragraphs[1]
    replace_text(header, "Тема проекта: ____________________________________________", "Тема проекта: Сортировка в современных языках программирования")
    replace_text(header, "Номер команды: _______", "Номер команды: 7")
    replace_text(header, "Лидер команды: ____________________________________________", "Лидер команды: Подъячий Глеб Русланович")
    replace_text(doc.paragraphs[4], "ФИО: ____________________________________________", "ФИО: " + person["name"])
    for paragraph in doc.paragraphs:
        for role in person["roles"]:
            needle = "☐ " + role
            if needle in paragraph.text:
                replace_text(paragraph, needle, "☑ " + role)
    for index, (task, result) in enumerate(person["tasks"], 1):
        set_cell(doc.tables[0].rows[index].cells[0], task)
        set_cell(doc.tables[0].rows[index].cells[1], result)
    for index, score in enumerate(("2", "2", "1", "1", "6"), 1):
        set_cell(doc.tables[1].rows[index].cells[1], score)
    for index, (name, what) in enumerate(PEERS, 1):
        set_cell(doc.tables[2].rows[index].cells[0], name)
        set_cell(doc.tables[2].rows[index].cells[1], "5")
        set_cell(doc.tables[2].rows[index].cells[2], what)
    for index, score in enumerate(("2", "2", "1", "5"), 1):
        set_cell(doc.tables[3].rows[index].cells[1], score)
    replace_text(doc.paragraphs[18], "______", "20")
    doc.paragraphs[29].text = person["hard"]
    doc.paragraphs[31].text = person["best"]
    doc.paragraphs[33].text = person["better"]
    replace_text(doc.paragraphs[37], "Дата: _____________________", "Дата: 28.10.2026")
    doc.save(target)
    print(target.name)


if __name__ == "__main__":
    for person in PEOPLE:
        fill(person)
